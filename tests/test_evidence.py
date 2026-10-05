"""カスタム証跡画像：画像ユーティリティ（Pillow 不要）、画像 API、ライブラリ設定の保存、Excel への挿入。"""
import http.client
import json
import os
import re
import shutil
import struct
import tempfile
import unittest
import zipfile
import zlib
from urllib.parse import quote

import _util
from openpyxl import load_workbook


def make_png(w, h, rgb=(35, 47, 62)):
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def make_jpeg_header(w, h):
    # SOI + APP0(JFIF) + SOF0 だけの最小ヘッダ（サイズ読み取りのテスト用）
    app0 = b"\xff\xe0" + struct.pack(">H", 16) + b"JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    sof = b"\xff\xc0" + struct.pack(">HBHHB", 11, 8, h, w, 1) + b"\x01\x11\x00"
    return b"\xff\xd8" + app0 + sof + b"\xff\xd9"


class ImagesUtilTest(unittest.TestCase):
    def test_sniff(self):
        from ba.images import sniff
        self.assertEqual(sniff(make_png(40, 30)), ("png", 40, 30))
        self.assertEqual(sniff(make_jpeg_header(1920, 1080)), ("jpeg", 1920, 1080))
        for bad in (b"", b"GIF89a....", b"<svg/>", b"\xff\xd8\xff"):
            with self.assertRaises(ValueError):
                sniff(bad)

    def test_fit(self):
        from ba.images import fit
        self.assertEqual(fit(1920, 1080, 760, 900), (760, 428))
        self.assertEqual(fit(400, 2000, 760, 900), (180, 900))
        self.assertEqual(fit(100, 50, 760, 900), (100, 50))        # 拡大はしない

    def test_raw_image_without_pillow(self):
        from ba.images import raw_image
        data = make_png(10, 10)
        img = raw_image(data, 10, 10)
        self.assertEqual((img.width, img.height, img.format), (10, 10, "png"))
        self.assertEqual(img._data(), data)


class FillEvidenceTest(unittest.TestCase):
    def test_sheet_and_mapped_cell(self):
        from ba.fill import fill_template
        tmp = tempfile.mkdtemp()
        try:
            p1, p2 = os.path.join(tmp, "a.png"), os.path.join(tmp, "b.png")
            with open(p1, "wb") as f:
                f.write(make_png(1280, 720))
            with open(p2, "wb") as f:
                f.write(make_png(300, 200, (200, 30, 30)))
            items = [{"id": "構築結果!E9", "sheet": "構築結果", "cell": "E9", "label": "画面", "source": {"kind": "evidence", "key": "img_ec2"}},
                     {"id": "構築結果!C9", "sheet": "構築結果", "cell": "C9", "label": "サーバ名", "type": "text", "source": {"kind": "fixed", "value": "web01"}}]
            ev = [(3, "EC2 を確認する", "EC2 詳細画面のスクリーンショット", "img_ec2", p1, {"w": 1280, "h": 720}),
                  (3, "EC2 を確認する", "", "", p2, {"w": 300, "h": 200})]
            out = os.path.join(tmp, "out.xlsx")
            written = fill_template(_util.TEMPLATE, items, {"構築結果!C9": "web01"}, out, evidence=ev)
            self.assertIn("構築結果!E9", written)
            with zipfile.ZipFile(out) as z:       # 閉じないと Windows では一時フォルダを削除できない
                media = [n for n in z.namelist() if n.startswith("xl/media/")]
                self.assertEqual(len(media), 3)
                self.assertTrue(all(z.read(m)[:4] == b"\x89PNG" for m in media))
                drawings = {n: z.read(n).decode("utf-8") for n in z.namelist() if re.match(r"xl/drawings/drawing\d+\.xml$", n)}
                self.assertEqual(len(drawings), 2)
                self.assertTrue(any(re.search(r"<(xdr:)?col>4</(xdr:)?col><(xdr:)?colOff>0</(xdr:)?colOff><(xdr:)?row>8<", x) for x in drawings.values()))
                for n in drawings:
                    rels = z.read(n.replace("drawings/", "drawings/_rels/") + ".rels").decode("utf-8")
                    self.assertIn("relationships/image", rels)
            wb = load_workbook(out)
            self.assertEqual(wb.sheetnames[-1], "証跡")
            ws = wb["証跡"]
            self.assertEqual(ws["A3"].value, "手順 3：EC2 を確認する")
            self.assertEqual(ws["A4"].value, "EC2 詳細画面のスクリーンショット　[img_ec2]")
            self.assertEqual(ws["A5"].value, None)
            self.assertEqual(wb["構築結果"]["C9"].value, "web01")
            self.assertIsNone(wb["構築結果"]["E9"].value)
        finally:
            shutil.rmtree(tmp)

    def test_no_evidence_is_unchanged(self):
        from ba.fill import fill_template
        tmp = tempfile.mkdtemp()
        try:
            out = os.path.join(tmp, "out.xlsx")
            fill_template(_util.TEMPLATE, [], {}, out, evidence=[])
            self.assertNotIn("証跡", load_workbook(out).sheetnames)
            self.assertFalse([n for n in zipfile.ZipFile(out).namelist() if n.startswith("xl/media/")])
        finally:
            shutil.rmtree(tmp)


class EvidenceApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from ba.server import App
        cls.tmp = tempfile.mkdtemp()
        cls.app = App(cls.tmp, port=0)
        cls.app.start_background()

    @classmethod
    def tearDownClass(cls):
        cls.app.httpd.shutdown()
        cls.app.httpd.server_close()
        shutil.rmtree(cls.tmp)

    def req(self, method, path, body=None, headers=None, token=True):
        c = http.client.HTTPConnection("127.0.0.1", self.app.port, timeout=20)
        h = dict(headers or {})
        if token:
            h["X-Token"] = self.app.token
        if isinstance(body, (dict, list)):
            body = json.dumps(body).encode("utf-8")
            h["Content-Type"] = "application/json"
        c.request(method, path, body=body, headers=h)
        r = c.getresponse()
        data = r.read()
        c.close()
        ct = r.getheader("Content-Type", "")
        return r.status, (json.loads(data.decode("utf-8")) if data and ct.startswith("application/json") else data), r

    def test_image_api(self):
        png = make_png(64, 48)
        st, m, _ = self.req("POST", "/api/sop/images", png, {"Content-Type": "image/png"})
        self.assertEqual(st, 200, m)
        self.assertTrue(re.match(r"^img-\d{14}-[0-9a-f]{8}$", m["id"]))
        self.assertEqual((m["ext"], m["w"], m["h"], m["size"]), ("png", 64, 48, len(png)))
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "sop_images", m["id"] + ".png")))
        st, data, r = self.req("GET", "/api/sop/images/" + m["id"])
        self.assertEqual((st, data, r.getheader("Content-Type")), (200, png, "image/png"))
        st, data, _ = self.req("GET", "/api/sop/images/%s?t=%s" % (m["id"], self.app.token), token=False)   # <img src> 用
        self.assertEqual(st, 200)
        self.assertEqual(self.req("GET", "/api/sop/images/" + m["id"], token=False)[0], 403)
        self.assertEqual(self.req("POST", "/api/sop/images", b"not an image")[0], 400)
        self.assertEqual(self.req("POST", "/api/sop/images", png, token=False)[0], 403)
        for bad in ("../../app", "img-1-x", "..%2f..%2fapp.py"):
            self.assertEqual(self.req("GET", "/api/sop/images/" + bad)[0], 404)
        st, j, _ = self.req("POST", "/api/sop/images", make_jpeg_header(800, 600))
        self.assertEqual((st, j["ext"], j["mime"], j["w"], j["h"]), (200, "jpg", "image/jpeg", 800, 600))
        self.assertEqual(self.req("GET", "/api/sop/images/" + j["id"])[2].getheader("Content-Type"), "image/jpeg")
        self.assertEqual(self.req("DELETE", "/api/sop/images/" + m["id"])[0], 200)
        self.assertEqual(self.req("GET", "/api/sop/images/" + m["id"])[0], 404)

    def test_library_settings_and_deliverable(self):
        with open(_util.DOCX, "rb") as f:
            st, pr, _ = self.req("POST", "/api/library/upload?type=procedure", f.read(), {"X-Filename": quote(os.path.basename(_util.DOCX))})
        ev = {"docHash": "abc", "docName": "x.docx", "steps": [{"index": 2, "title": "踏み台へ接続", "items": [{"id": "ev1", "desc": "EC2 詳細画面のスクリーンショット", "key": "img_ec2"}]}]}
        st, m, _ = self.req("PUT", "/api/library/" + pr["id"], {"evidence": ev})
        self.assertEqual(m["evidence"], ev)
        st, m, _ = self.req("PUT", "/api/library/" + pr["id"], {"name": "改名"})      # 他の更新で消えない
        st, lst, _ = self.req("GET", "/api/library?type=procedure")
        self.assertEqual(lst["items"][0]["evidence"], ev)
        # 手順実行セッション（画像メタのみ）→ 成果物に ?sop= で挿入
        with open(_util.TEMPLATE, "rb") as f:
            st, xt, _ = self.req("POST", "/api/library/upload?type=excel_template", f.read(), {"X-Filename": quote("tpl.xlsx")})
        st, img, _ = self.req("POST", "/api/sop/images", make_png(200, 100))
        key = "sopRunner:v1:x.docx:abc"
        sess = {"docName": "x.docx", "steps": [{"id": "s1", "title": "A"}, {"id": "s2", "title": "B", "evidence": ev["steps"][0]["items"]}],
                "results": {"s2": {"images": {"ev1": [img]}}}}
        self.assertEqual(self.req("PUT", "/api/sop/sessions/" + quote(key, safe=""), sess)[0], 200)
        st, job, _ = self.req("POST", "/api/jobs", {"name": "t", "template_id": xt["id"]})
        for url, n_media in (("/api/jobs/%s/deliverable.xlsx" % job["id"], 0),
                             ("/api/jobs/%s/deliverable.xlsx?sop=%s" % (job["id"], quote(key, safe="")), 1)):
            st, data, _ = self.req("GET", url)
            self.assertEqual(st, 200)
            p = os.path.join(self.tmp, "d.xlsx")
            with open(p, "wb") as f:
                f.write(data)
            self.assertEqual(len([n for n in zipfile.ZipFile(p).namelist() if n.startswith("xl/media/")]), n_media)
            self.assertEqual("証跡" in load_workbook(p).sheetnames, bool(n_media))


if __name__ == "__main__":
    unittest.main()
