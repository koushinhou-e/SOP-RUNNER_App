import os
import shutil
import tempfile
import unittest

import _util
from ba import fill, paramsheet, xlsx_detect
from openpyxl import load_workbook


def style_sig(c):
    return (c.font.b, c.font.sz, c.font.color.rgb if c.font.color is not None else None, c.fill.fill_type,
            c.fill.fgColor.rgb, c.border.left.style, c.border.right.style, c.border.top.style, c.border.bottom.style,
            c.number_format, c.alignment.horizontal)


class FillTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.out = os.path.join(self.tmp, "out.xlsx")
        items = xlsx_detect.detect_workbook(_util.TEMPLATE)
        ps = paramsheet.parse(_util.PARAMS)
        xlsx_detect.auto_map(items, ps["params"])
        ctx = paramsheet.server_context(ps, "web02")
        values = {}
        for it in items:
            if it["source"]["kind"] == "param":
                values[it["id"]] = ctx[it["source"]["key"]]
        by = {i["cell"]: i for i in items}
        values[by["C3"]["id"]] = "サンプル基盤構築"
        values[by["C4"]["id"]] = "2026/10/03"
        values[by["C5"]["id"]] = "山田 太郎"
        values[by["C6"]["id"]] = "佐藤 花子"
        values[by["D12"]["id"]] = "2"            # 確認結果（作業結果）の数値は数値として書く
        values[by["D15"]["id"]] = "007"          # 先頭ゼロ付きは文字列のまま
        values[by["E11"]["id"]] = "OK"
        values[by["C21"]["id"]] = "合格"
        self.items = items
        self.written = fill.fill_template(_util.TEMPLATE, items, values, self.out)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_values_written_with_types(self):
        ws = load_workbook(self.out)["構築結果"]
        self.assertEqual(ws["C9"].value, "web02.example.local")
        self.assertEqual(ws["C11"].value, "t3.large")
        self.assertEqual(ws["C12"].value, 2)               # number item → int
        self.assertEqual(ws["C13"].value, 8)
        self.assertEqual(ws["C4"].value.strftime("%Y-%m-%d"), "2026-10-03")   # date-formatted cell → datetime
        self.assertEqual(ws["C3"].value, "サンプル基盤構築")   # {{project_name}} replaced
        self.assertEqual(ws["C6"].value, "佐藤 花子")
        self.assertEqual(ws["E11"].value, "OK")
        self.assertEqual(ws["D12"].value, 2)
        self.assertEqual(ws["D15"].value, "007")
        self.assertEqual(ws["C21"].value, "合格")
        self.assertIsNone(ws["F9"].value)                  # "不填" cells untouched
        self.assertIsNone(ws["D9"].value)
        self.assertEqual(ws["C22"].value, "（　　　）")       # not filled → placeholder kept for hand edit
        self.assertEqual(len(self.written), 19)

    def test_formatting_preserved(self):
        a, b = load_workbook(_util.TEMPLATE), load_workbook(self.out)
        self.assertEqual(a.sheetnames, b.sheetnames)
        self.assertEqual(b["リスト"].sheet_state, "hidden")
        for name in a.sheetnames:
            wa, wb_ = a[name], b[name]
            self.assertEqual(sorted(map(str, wa.merged_cells.ranges)), sorted(map(str, wb_.merged_cells.ranges)))
            self.assertEqual({k: v.width for k, v in wa.column_dimensions.items()}, {k: v.width for k, v in wb_.column_dimensions.items()})
            self.assertEqual({k: v.height for k, v in wa.row_dimensions.items()}, {k: v.height for k, v in wb_.row_dimensions.items()})
            self.assertEqual(wa.freeze_panes, wb_.freeze_panes)
            dva = sorted((d.type, d.formula1, str(d.sqref)) for d in wa.data_validations.dataValidation)
            dvb = sorted((d.type, d.formula1, str(d.sqref)) for d in wb_.data_validations.dataValidation)
            self.assertEqual(dva, dvb)
            for row in wa.iter_rows():
                for c in row:
                    self.assertEqual(style_sig(c), style_sig(wb_[c.coordinate]), "%s!%s" % (name, c.coordinate))
            for row in wa.iter_rows():                 # untouched cells keep their values
                for c in row:
                    if "%s!%s" % (name, c.coordinate) not in self.written:
                        self.assertEqual(c.value, wb_[c.coordinate].value)
