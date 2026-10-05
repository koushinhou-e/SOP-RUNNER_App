import unittest

import _util
from ba import paramsheet, xlsx_detect


class DetectTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.items = xlsx_detect.detect_workbook(_util.TEMPLATE)
        cls.by = {i["cell"]: i for i in cls.items}

    def test_list_validation_whole_column(self):
        # 入力規則のリストが列全体（=選択肢!$A:$A）を参照していても検出できること（以前は TypeError で取り込み失敗）
        import os
        import tempfile
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = Workbook()
        ws = wb.active
        ws.title = "報告書"
        ws["A1"], ws["A2"] = "OS", "リージョン"
        opts = wb.create_sheet("選択肢")
        for i, v in enumerate(["Amazon Linux 2023", "RHEL 9"], 1):
            opts.cell(row=i, column=1, value=v)
        dv = DataValidation(type="list", formula1="'選択肢'!$A:$A")
        ws.add_data_validation(dv)
        dv.add("B1")
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            wb.save(path)
            items = {i["id"]: i for i in xlsx_detect.detect_workbook(path)}
        finally:
            os.remove(path)
        self.assertEqual(items["報告書!B1"]["type"], "dropdown")
        self.assertEqual(items["報告書!B1"]["options"], ["Amazon Linux 2023", "RHEL 9"])

    def test_counts_and_reasons(self):
        self.assertEqual(len(self.items), 50)
        reasons = {}
        for i in self.items:
            reasons[i["reason"]] = reasons.get(i["reason"], 0) + 1
        self.assertEqual(reasons, {"placeholder": 3, "highlight": 2, "label": 32, "validation": 13})

    def test_placeholders(self):
        self.assertEqual(self.by["C3"]["key"], "project_name")
        self.assertEqual(self.by["C3"]["label"], "案件名")
        self.assertEqual(self.by["C6"]["placeholder"], "＿＿＿＿＿＿")
        self.assertEqual(self.by["C22"]["label"], "特記事項")
        self.assertFalse(self.by["C22"]["rules"]["required"])

    def test_highlight_and_date(self):
        self.assertEqual(self.by["C4"]["reason"], "highlight")
        self.assertEqual(self.by["C4"]["type"], "date")
        self.assertEqual(self.by["C5"]["label"], "作業者")
        self.assertEqual(self.by["C5"]["key"], "worker")

    def test_validation_list_inline_and_range(self):
        e9 = self.by["E9"]
        self.assertEqual((e9["type"], e9["options"], e9["label"]), ("dropdown", ["OK", "NG", "対象外"], "サーバ名 / 判定"))
        d9 = self.by["D9"]                      # 確認結果の列は自由記述（作業中に入力した値を記入する）
        self.assertEqual((d9["type"], d9["label"], d9["col_header"]), ("text", "サーバ名 / 確認結果", "確認結果"))
        self.assertEqual(self.by["C21"]["options"], ["合格", "不合格"])   # =リスト!$A$1:$A$2 (hidden sheet)

    def test_number_types(self):
        self.assertEqual(self.by["C12"]["type"], "number")
        self.assertEqual((self.by["C12"]["rules"]["min"], self.by["C12"]["rules"]["max"]), (1, 128))
        self.assertEqual(self.by["C13"]["type"], "number")   # number_format "0"
        self.assertTrue(self.by["C13"]["rules"]["required"])  # メモリ is not メモ

    def test_presets_by_label(self):
        self.assertEqual(self.by["C10"]["rules"].get("preset"), "instance_id")
        self.assertEqual(self.by["C18"]["rules"].get("preset"), "ipv4")
        self.assertEqual(self.by["C9"]["rules"].get("preset"), "hostname")

    def test_notes_column_optional_and_unfilled(self):
        f9 = self.by["F9"]
        self.assertEqual(f9["source"]["kind"], "none")
        self.assertFalse(f9["rules"]["required"])

    def test_merged_cells_not_duplicated(self):
        self.assertNotIn("D21", self.by)       # C21:D21 merged
        self.assertNotIn("C1", self.by)        # title merged B1:F1

    def test_automap(self):
        items = xlsx_detect.detect_workbook(_util.TEMPLATE)
        ps = paramsheet.parse(_util.PARAMS)
        n = xlsx_detect.auto_map(items, ps["params"])
        self.assertEqual(n, 33)
        by = {i["cell"]: i for i in items}
        self.assertEqual(by["C11"]["source"], {"kind": "param", "key": "instance_type"})
        self.assertEqual(by["C13"]["source"], {"kind": "param", "key": "memory_gib"})
        self.assertEqual(by["D11"]["source"], {"kind": "result", "key": "instance_type"})   # 確認結果 = 作業中に入力した値（パラメータ値ではない）
        self.assertEqual(by["E11"]["source"], {"kind": "judge", "key": "instance_type"})    # 判定 = 設定値との一致で OK/NG
        self.assertEqual(by["F11"]["source"]["kind"], "none")                                 # 備考は記入しない
        self.assertEqual(by["C21"]["source"]["kind"], "input")                                # 総合判定・案件名などは作業入力のまま
        self.assertEqual(by["C3"]["source"]["kind"], "input")

    def test_automap_other_layouts(self):
        """設定値 / 実績値 の様式、確認結果が OK/NG リストの様式（以前のサンプル）でも適切に割り当てる。"""
        params = [{"key": "instance_type", "label": "インスタンスタイプ"}, {"key": "vcpu", "label": "vCPU"}]

        def item(cell, row, header, typ="text", options=None):
            return {"id": "S!" + cell, "cell": cell, "row_label": row, "col_header": header, "label": row + (" / " + header if header else ""),
                    "type": typ, "options": options or [], "source": {"kind": "input"}}
        items = [item("C2", "インスタンスタイプ", None), item("D2", "インスタンスタイプ", "実績値"), item("E2", "インスタンスタイプ", "確認結果", "dropdown", ["○", "×"]),
                 item("F2", "インスタンスタイプ", "備考"), item("D3", "vCPU", "合否", "dropdown", ["合格", "不合格"]), item("D4", "作業者", "確認結果")]
        self.assertEqual(xlsx_detect.auto_map(items, params), 4)
        kinds = {i["cell"]: (i["source"]["kind"], i["source"].get("key")) for i in items}
        self.assertEqual(kinds, {"C2": ("param", "instance_type"), "D2": ("result", "instance_type"), "E2": ("judge", "instance_type"),
                                 "F2": ("input", None), "D3": ("judge", "vcpu"), "D4": ("input", None)})

    def test_judge_marks(self):
        self.assertEqual(xlsx_detect.judge_marks(["OK", "NG", "対象外"]), ("OK", "NG"))
        self.assertEqual(xlsx_detect.judge_marks(["ok", "ng"]), ("ok", "ng"))
        self.assertEqual(xlsx_detect.judge_marks(["合格", "不合格"]), ("合格", "不合格"))
        self.assertEqual(xlsx_detect.judge_marks(["○", "×"]), ("○", "×"))
        self.assertEqual(xlsx_detect.judge_marks([]), ("OK", "NG"))


class ParamSheetTest(unittest.TestCase):
    def test_parse(self):
        p = paramsheet.parse(_util.PARAMS)
        self.assertEqual(p["servers"], ["web01", "web02", "db01"])
        self.assertEqual(len(p["params"]), 13)
        ctx = paramsheet.server_context(p, "db01")
        self.assertEqual(ctx["instance_type"], "r5.xlarge")
        self.assertEqual(ctx["memory_gib"], "32")      # int, not "32.0"
        self.assertEqual(ctx["private_ip"], "198.51.100.21")
        self.assertEqual([x["category"] for x in p["params"]][:2], ["基本", "基本"])

    def test_two_column_sheet(self):
        import os
        import tempfile
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["項目", "値"])
        ws.append(["インスタンスタイプ", "t3.small"])
        ws.append(["EBSサイズ", 50])
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        try:
            wb.save(path)
            p = paramsheet.parse(path)
            self.assertEqual(p["servers"], ["default"])
            self.assertEqual(p["params"][1]["values"]["default"], "50")
            self.assertEqual(p["params"][0]["key"], "インスタンスタイプ")
        finally:
            os.remove(path)
