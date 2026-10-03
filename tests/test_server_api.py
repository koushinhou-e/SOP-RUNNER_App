"""HTTP API 全流程 + 安全检查 + “无外部网络访问”断言。"""
import http.client
import json
import os
import shutil
import socket
import tempfile
import unittest
from urllib.parse import quote

import _util
from openpyxl import load_workbook

ATTEMPTS = []
_orig = {"connect": socket.socket.connect, "connect_ex": socket.socket.connect_ex, "create_connection": socket.create_connection, "getaddrinfo": socket.getaddrinfo}
LOOPBACK = ("127.0.0.1", "localhost", "::1")


def _host(addr):
    return addr[0] if isinstance(addr, tuple) else str(addr)


def _guard_connect(self, addr):
    if self.family in (socket.AF_INET, socket.AF_INET6) and _host(addr) not in LOOPBACK:
        ATTEMPTS.append(("connect", addr))
        raise OSError("outbound network blocked by test")
    return _orig["connect"](self, addr)


def _guard_connect_ex(self, addr):
    if self.family in (socket.AF_INET, socket.AF_INET6) and _host(addr) not in LOOPBACK:
        ATTEMPTS.append(("connect_ex", addr))
        return 111
    return _orig["connect_ex"](self, addr)


def _guard_gai(host, *a, **k):
    if host not in LOOPBACK and host is not None:
        ATTEMPTS.append(("getaddrinfo", host))
        raise socket.gaierror("DNS blocked by test")
    return _orig["getaddrinfo"](host, *a, **k)


class ApiFlowTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        socket.socket.connect = _guard_connect
        socket.socket.connect_ex = _guard_connect_ex
        socket.getaddrinfo = _guard_gai
        from ba.server import App
        cls.tmp = tempfile.mkdtemp()
        cls.app = App(cls.tmp, port=0)
        cls.app.start_background()
        cls.port = cls.app.port

    @classmethod
    def tearDownClass(cls):
        cls.app.httpd.shutdown()
        cls.app.httpd.server_close()
        shutil.rmtree(cls.tmp)
        socket.socket.connect = _orig["connect"]
        socket.socket.connect_ex = _orig["connect_ex"]
        socket.getaddrinfo = _orig["getaddrinfo"]

    def req(self, method, path, body=None, headers=None, token=True, raw=False):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=20)
        h = dict(headers or {})
        if token:
            h["X-Token"] = self.app.token
        if isinstance(body, (dict, list)):
            body = json.dumps(body).encode("utf-8")
            h["Content-Type"] = "application/json"
        c.request(method, path, body=body, headers=h)
        r = c.getresponse()
        data = r.read()
        c.close()
        if raw:
            return r, data
        return r.status, (json.loads(data.decode("utf-8")) if data and r.getheader("Content-Type", "").startswith("application/json") else data)

    def upload(self, typ, path):
        with open(path, "rb") as f:
            return self.req("POST", "/api/library/upload?type=" + typ, f.read(), {"X-Filename": quote(os.path.basename(path))})

    def test_1_security(self):
        self.assertEqual(self.app.httpd.server_address[0], "127.0.0.1")
        st, _ = self.req("GET", "/api/library", token=False)
        self.assertEqual(st, 403)
        st, _ = self.req("GET", "/api/library", headers={"X-Token": "wrong"}, token=False)
        self.assertEqual(st, 403)
        st, _ = self.req("GET", "/", headers={"Host": "attacker.example:%d" % self.port}, token=False)
        self.assertEqual(st, 403)
        st, _ = self.req("GET", "/../ba/server.py", token=False)
        self.assertEqual(st, 404)
        st, _ = self.req("GET", "/%2e%2e/app.py", token=False)
        self.assertEqual(st, 404)
        st, _ = self.req("GET", "/api/library/..%2f..%2fetc", token=True)
        self.assertEqual(st, 404)
        r, html = self.req("GET", "/", token=False, raw=True)
        self.assertEqual(r.status, 200)
        self.assertIn(self.app.token.encode(), html)
        self.assertIn("default-src 'self'", r.getheader("Content-Security-Policy"))

    def test_2_full_flow(self):
        st, ps = self.upload("param_sheet", _util.PARAMS)
        self.assertEqual(st, 200, ps)
        st, xt = self.upload("excel_template", _util.TEMPLATE)
        self.assertEqual((st, len(xt["items"])), (200, 39))
        st, r = self.req("POST", "/api/library/%s/automap" % xt["id"], {"param_sheet_id": ps["id"]})
        self.assertEqual(r["mapped"], 11)
        st, pr = self.upload("procedure", _util.DOCX)
        self.assertEqual(st, 200)
        with open(_util.CMDS, encoding="utf-8") as f:
            cs_body = json.load(f)
        st, cs = self.req("POST", "/api/library/command_set", cs_body)
        self.assertEqual(len(cs["templates"]), 9)
        # rename + list
        st, m = self.req("PUT", "/api/library/" + pr["id"], {"name": "パッチ手順書（改名）"})
        self.assertEqual(m["name"], "パッチ手順書（改名）")
        st, lst = self.req("GET", "/api/library?type=procedure")
        self.assertEqual([x["name"] for x in lst["items"]], ["パッチ手順書（改名）"])
        # file download round trip
        r, data = self.req("GET", "/api/library/%s/file?t=%s" % (pr["id"], self.app.token), token=False, raw=True)
        with open(_util.DOCX, "rb") as f:
            self.assertEqual(data, f.read())
        # job
        st, job = self.req("POST", "/api/jobs", {"name": "web01 構築確認", "template_id": xt["id"], "param_sheet_id": ps["id"], "server": "web01", "command_set_id": cs["id"]})
        jid = job["id"]
        st, v = self.req("GET", "/api/jobs/" + jid)
        res = {r_["cell"]: r_ for r_ in v["resolved"]}
        self.assertEqual(res["C11"]["value"], "t3.large")
        self.assertEqual(res["C5"]["errors"][0]["code"], "required")
        items = {i["cell"]: i["id"] for i in xt["items"]}
        st, v = self.req("PUT", "/api/jobs/" + jid, {"inputs": {items["C3"]: "サンプル案件", items["C4"]: "2026/10/03", items["C5"]: "山田 太郎", items["C6"]: "佐藤 花子", items["D11"]: "OK", items["C21"]: "Maybe"}})
        res = {r_["cell"]: r_ for r_ in v["resolved"]}
        self.assertEqual(res["C5"]["errors"], [])
        self.assertEqual(res["C21"]["errors"][0]["code"], "options")
        # commands + scripts
        st, c = self.req("GET", "/api/jobs/%s/commands" % jid)
        self.assertEqual(len(c["commands"]), 9)
        self.assertIn("i-0123456789abcdef0", c["commands"][0]["sh"])
        r, sh = self.req("GET", "/api/jobs/%s/commands.sh?t=%s" % (jid, self.app.token), token=False, raw=True)
        self.assertTrue(sh.startswith(b"#!/usr/bin/env bash"))
        self.assertIn("attachment", r.getheader("Content-Disposition"))
        r, ps1 = self.req("GET", "/api/jobs/%s/commands.ps1?t=%s" % (jid, self.app.token), token=False, raw=True)
        self.assertTrue(ps1.startswith(b"\xef\xbb\xbf# "))
        self.assertIn(b"\r\n", ps1)
        # globals override
        st, v = self.req("PUT", "/api/jobs/" + jid, {"globals": {"region": "us-east-1"}})
        st, c = self.req("GET", "/api/jobs/%s/commands" % jid)
        self.assertIn("--region us-east-1", c["commands"][0]["sh"])
        # compare
        st, v = self.req("PUT", "/api/jobs/" + jid, {"compare": {"instance_type": {"actual": "t3.medium", "judgement": "NG", "note": "要確認"},
                                                                 "memory_gib": {"actual": "8 GiB", "judgement": "OK"}}})
        st, cmp_ = self.req("GET", "/api/jobs/%s/compare" % jid)
        rows = {x["key"]: x for x in cmp_["rows"]}
        self.assertEqual(rows["instance_type"]["auto"], "mismatch")
        self.assertEqual(rows["memory_gib"]["auto"], "match")
        self.assertTrue(rows["instance_type"]["commands"])
        self.assertEqual(cmp_["summary"]["ng"], 1)
        r, data = self.req("GET", "/api/jobs/%s/compare.xlsx?t=%s" % (jid, self.app.token), token=False, raw=True)
        p = os.path.join(self.tmp, "cmp.xlsx")
        with open(p, "wb") as f:
            f.write(data)
        ws = load_workbook(p).active
        vals = [[c.value for c in row] for row in ws.iter_rows()]
        flat = [x for row in vals for x in row]
        self.assertIn("t3.medium", flat)
        self.assertIn("不一致", flat)
        self.assertIn("NG", flat)
        # deliverable
        r, data = self.req("GET", "/api/jobs/%s/deliverable.xlsx?t=%s" % (jid, self.app.token), token=False, raw=True)
        self.assertEqual(r.status, 200)
        p = os.path.join(self.tmp, "deliv.xlsx")
        with open(p, "wb") as f:
            f.write(data)
        ws = load_workbook(p)["構築結果"]
        self.assertEqual((ws["C9"].value, ws["C10"].value, ws["C5"].value, ws["D11"].value), ("web01.example.local", "i-0123456789abcdef0", "山田 太郎", "OK"))
        # SOP sessions server-side
        st, _ = self.req("PUT", "/api/sop/sessions/" + quote("sopRunner:v1:a.docx:abc", safe=""), {"steps": [], "docName": "a.docx"})
        st, ss = self.req("GET", "/api/sop/sessions")
        self.assertEqual(ss["items"][0]["key"], "sopRunner:v1:a.docx:abc")
        st, _ = self.req("DELETE", "/api/sop/sessions/" + quote("sopRunner:v1:a.docx:abc", safe=""))
        st, ss = self.req("GET", "/api/sop/sessions")
        self.assertEqual(ss["items"], [])
        # delete
        st, _ = self.req("DELETE", "/api/library/" + pr["id"])
        st, lst = self.req("GET", "/api/library?type=procedure")
        self.assertEqual(lst["items"], [])
        self.assertTrue(os.path.isfile(os.path.join(self.tmp, "jobs", jid + ".json")))

    def test_3_bad_uploads(self):
        st, r = self.req("POST", "/api/library/upload?type=excel_template", b"not a zip", {"X-Filename": "x.xlsx"})
        self.assertEqual(st, 400)
        st, r = self.req("POST", "/api/library/upload?type=procedure", b"not a zip", {"X-Filename": "x.docx"})
        self.assertEqual(st, 400)
        st, lst = self.req("GET", "/api/library?type=excel_template")
        self.assertTrue(all(x.get("items") for x in lst["items"]))   # failed uploads are cleaned up

    def test_4_samples_endpoint(self):
        st, r = self.req("POST", "/api/samples")
        self.assertEqual((st, len(r["items"])), (200, 4))

    def test_9_no_outbound_network(self):
        self.assertEqual(ATTEMPTS, [], "outbound network attempts: %r" % ATTEMPTS)


class StaticOfflineTest(unittest.TestCase):
    ALLOWED_URL_PREFIXES = ("http://schemas.openxmlformats.org/", "http://purl.org/dc/", "http://www.w3.org/2001/XMLSchema-instance",
                            "http://stuartk.com/jszip", "https://github.com/nodeca/pako", "https://raw.github.com/Stuk/jszip", "https://stuk.github.io/jszip")

    def test_web_has_no_external_resources(self):
        import re
        web = os.path.join(_util.ROOT, "web")
        bad = []
        for dp, _, fns in os.walk(web):
            for fn in fns:
                if not fn.endswith((".html", ".js", ".css")):
                    continue
                with open(os.path.join(dp, fn), encoding="utf-8") as f:
                    s = f.read()
                for u in re.findall(r"https?://[^\s\"'<>)`]+", s):
                    if not u.startswith(self.ALLOWED_URL_PREFIXES):
                        bad.append((fn, u))
                for pat in (r"<script[^>]+src=[\"']?https?:", r"<link[^>]+href=[\"']?https?:", r"url\(\s*[\"']?https?:", r"@import", r"\bWebSocket\b", r"sendBeacon", r"EventSource"):
                    if re.search(pat, s):
                        bad.append((fn, pat))
        self.assertEqual(bad, [])

    def test_python_has_no_network_clients(self):
        import re
        bad = []
        for dp, _, fns in os.walk(os.path.join(_util.ROOT, "ba")):
            for fn in fns:
                if fn.endswith(".py"):
                    with open(os.path.join(dp, fn), encoding="utf-8") as f:
                        s = f.read()
                    for pat in (r"urllib\.request", r"http\.client", r"\brequests\b", r"ftplib", r"smtplib", r"socket\.create_connection", r"urlopen"):
                        if re.search(pat, s):
                            bad.append((fn, pat))
        self.assertEqual(bad, [])


if __name__ == "__main__":
    unittest.main()
