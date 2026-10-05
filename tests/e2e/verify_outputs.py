"""校验 E2E 下载的文件（用 openpyxl）：交付物的值与格式、比对结果、证迹图片（docx / xlsx 的 zip 结构 + LibreOffice 渲染）。
  python tests/e2e/verify_outputs.py"""
import os
import re
import shutil
import subprocess
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path[:0] = [os.path.join(ROOT, "vendor"), os.path.join(ROOT, "tests")]
from openpyxl import load_workbook  # noqa: E402
from test_fill import style_sig  # noqa: E402

OUT = os.path.join(ROOT, "tests", "_out")
TPL = os.path.join(ROOT, "samples", "構築結果報告書_template_sample.xlsx")
ok = bad = 0


def check(c, m):
    global ok, bad
    print(("  ✔ " if c else "  ✘ ") + m)
    ok += bool(c)
    bad += (not c)


d = load_workbook(os.path.join(OUT, "deliverable.xlsx"))
t = load_workbook(TPL)
ws, wt = d["構築結果"], t["構築結果"]
check(ws["C9"].value == "web01.example.local" and ws["C10"].value == "i-0123456789abcdef0" and ws["C11"].value == "t3.large", "交付物：参数值已写入")
check(ws["C12"].value == 2 and ws["C13"].value == 8 and ws["C19"].value == 30, "交付物：数字单元格写为数值")
check(ws["C4"].value.strftime("%Y-%m-%d") == "2026-10-03" and ws["C4"].number_format == "yyyy/mm/dd", "交付物：作業日写为日期并保留 yyyy/mm/dd 格式")
check(ws["C3"].value == "サンプル基盤構築" and ws["C5"].value == "山田 太郎" and ws["C6"].value == "佐藤 花子" and ws["C7"].value == "MNG-0042", "交付物：作业输入已写入（含手动新增的 C7）")
check(ws["D9"].value == "web01.example.local" and ws["D11"].value == "t3.medium" and ws["D12"].value == 2 and ws["D13"].value == "8 GiB", "交付物：確認結果 = ③ 的实测值（数字写为数值）")
check(ws["E11"].value == "NG" and all(ws["E%d" % r].value == "OK" for r in range(9, 20) if r != 11) and ws["C21"].value == "合格", "交付物：判定（与设定值一致 OK / 不一致 NG）/ 総合判定")
check(ws["F9"].value is None and ws["C22"].value == "（　　　）", "交付物：不填/未填的单元格保持原样（留给手动编辑）")
diffs = [c.coordinate for row in wt.iter_rows() for c in row if style_sig(c) != style_sig(ws[c.coordinate])]
check(not diffs, "交付物：所有单元格样式（字体/底色/边框/数字格式/对齐）与模板一致 %s" % (diffs[:5] if diffs else ""))
check(sorted(map(str, wt.merged_cells.ranges)) == sorted(map(str, ws.merged_cells.ranges)), "交付物：合并单元格一致")
check({k: v.width for k, v in wt.column_dimensions.items()} == {k: v.width for k, v in ws.column_dimensions.items()}, "交付物：列宽一致")
check(len(ws.data_validations.dataValidation) == len(wt.data_validations.dataValidation) and d["リスト"].sheet_state == "hidden", "交付物：数据验证与隐藏工作表保留")
check(ws.freeze_panes == wt.freeze_panes, "交付物：冻结窗格保留")

c = load_workbook(os.path.join(OUT, "compare.xlsx")).active
rows = [[x.value for x in r] for r in c.iter_rows()]
hdr_i = next(i for i, r in enumerate(rows) if r[0] == "No")
data = [r for r in rows[hdr_i + 1:] if r[0]]
it = next(r for r in data if r[3] == "instance_type")
check(len(data) == 13, "比对结果：13 行")
check(it[4] == "t3.large" and it[5] == "t3.medium" and it[6] == "不一致" and it[7] == "NG" and "変更要否" in (it[9] or ""), "比对结果：instance_type 不一致 / NG / 备注")
check(sum(1 for r in data if r[7] == "OK") == 12, "比对结果：OK 12")
check(c.cell(row=hdr_i + 1 + data.index(it) + 1, column=6).fill.fgColor.rgb.endswith("FDE2E2"), "比对结果：不一致行标红")

# ---- 证迹图片：Word 实施记录 ----
z = zipfile.ZipFile(os.path.join(OUT, "record.docx"))
names = z.namelist()
media = sorted(n for n in names if n.startswith("word/media/") and not n.endswith("/"))
doc = z.read("word/document.xml").decode("utf-8")
rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
ct = z.read("[Content_Types].xml").decode("utf-8")
check(len(media) == 2 and all(z.read(m)[:8] == b"\x89PNG\r\n\x1a\n" for m in media), "docx：word/media に PNG 2 枚")
embeds = re.findall(r'r:embed="([^"]+)"', doc)
check(len(embeds) == 2 and all(re.search(r'Id="%s"[^>]*relationships/image" Target="media/' % e, rels) for e in embeds), "docx：drawing の r:embed → image リレーション")
check('Extension="png" ContentType="image/png"' in ct, "docx：[Content_Types] に png")
ext = [int(x) for x in re.findall(r'<wp:extent cx="(\d+)"', doc)]
check(ext and max(ext) <= (14720 - 160) * 635, "docx：画像幅は表の内幅以下（%s EMU）" % ext)
check("証跡 3-1：EC2 詳細画面のスクリーンショット" in doc and "証跡 3-2：セキュリティグループ設定" in doc, "docx：各画像の説明キャプション")

# ---- 证迹图片：交付物 Excel ----
X = os.path.join(OUT, "deliverable_evidence.xlsx")
zx = zipfile.ZipFile(X)
xn = zx.namelist()
xmedia = [n for n in xn if n.startswith("xl/media/")]
drawings = [n for n in xn if re.match(r"xl/drawings/drawing\d+\.xml$", n)]
check(len(xmedia) == 3 and len(drawings) == 2, "xlsx：xl/media に画像 3 枚（証跡シート 2 + セル配置 1）、drawing 2 個 %s" % xmedia)
dw = {n: zx.read(n).decode("utf-8") for n in drawings}
anchored = [n for n, x in dw.items() if re.search(r"<(xdr:)?from><(xdr:)?col>5</(xdr:)?col><(xdr:)?colOff>0</(xdr:)?colOff><(xdr:)?row>8</", x)]
check(len(anchored) == 1, "xlsx：マッピング（構築結果!F9）の位置に画像を配置")
check("image/png" in zx.read("[Content_Types].xml").decode("utf-8") or 'Extension="png"' in zx.read("[Content_Types].xml").decode("utf-8"), "xlsx：Content_Types に png")
de = load_workbook(X)
check("証跡" in de.sheetnames and de["証跡"]["A3"].value.startswith("手順 3：") and "EC2 詳細画面のスクリーンショット" in de["証跡"]["A4"].value, "xlsx：「証跡」シートに手順名・説明")
dws = de["構築結果"]
check(dws["C9"].value == "web01.example.local" and not [c.coordinate for row in wt.iter_rows() for c in row if style_sig(c) != style_sig(dws[c.coordinate])], "xlsx：証跡あり出力でも値・書式は通常出力と同じ")

# ---- 確認結果報告書（手順実行の完了ページから出力）----
RP = os.path.join(OUT, "report.xlsx")
rws = load_workbook(RP)["構築結果"]
check([rws.cell(row=11, column=c).value for c in range(3, 6)] == ["t3.large", "t3.medium", "NG"] and rws["D9"].value == "web01.example.local" and rws["E9"].value == "OK",
      "確認結果報告書：設定値（パラメータシート）／確認結果（作業中の入力値）／判定")
with zipfile.ZipFile(RP) as zr:
    check(len([n for n in zr.namelist() if n.startswith("xl/media/")]) == 3, "確認結果報告書：証跡画像を挿入")

# ---- LibreOffice で PDF → PNG に描画（目視確認用、screenshots/ にコピー） ----
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
SHOT = os.path.join(ROOT, "screenshots")
if SOFFICE and shutil.which("pdftoppm"):
    R = os.path.join(OUT, "render")
    os.makedirs(R, exist_ok=True)
    for src, prefix, page in (("record.docx", "12-docx-render", 1), ("deliverable_evidence.xlsx", "13-xlsx-render", None)):
        subprocess.run([SOFFICE, "--headless", "-env:UserInstallation=file:///tmp/ba-lo-profile", "--convert-to", "pdf", "--outdir", R, os.path.join(OUT, src)],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        pdf = os.path.join(R, os.path.splitext(src)[0] + ".pdf")
        check(os.path.exists(pdf), "LibreOffice：%s → PDF" % src)
        if not os.path.exists(pdf):
            continue
        nimg = 0
        if shutil.which("pdfimages"):
            lst = subprocess.run(["pdfimages", "-list", pdf], stdout=subprocess.PIPE, universal_newlines=True).stdout.splitlines()[2:]
            nimg = len([l for l in lst if " image " in l or " smask " not in l])
            check(len([l for l in lst if l.split()[2:3] == ["image"]]) >= (2 if src.endswith("docx") else 3), "LibreOffice：%s の PDF に画像が描画されている（%d）" % (src, len(lst)))
        for f in os.listdir(R):
            if f.startswith(prefix):
                os.remove(os.path.join(R, f))
        subprocess.run(["pdftoppm", "-png", "-r", "60", pdf, os.path.join(R, prefix)], check=True)
        pages = sorted(f for f in os.listdir(R) if f.startswith(prefix) and f.endswith(".png"))
        for k, f in enumerate(pages):
            shutil.copy(os.path.join(R, f), os.path.join(SHOT, "%s-p%d.png" % (prefix, k + 1)))
        print("    render: %s → screenshots/%s-p1..%d.png" % (src, prefix, len(pages)))
else:
    print("  - soffice / pdftoppm が無いので描画確認をスキップ")

# ---- 作業入力値一覧 PDF ----
VP = os.path.join(OUT, "values.pdf")
check(os.path.exists(VP), "作業入力値一覧 PDF が出力されている")
if os.path.exists(VP) and shutil.which("pdfinfo") and shutil.which("pdftotext"):
    info = subprocess.run(["pdfinfo", VP], stdout=subprocess.PIPE, universal_newlines=True).stdout
    m = re.search(r"Page size:\s+([\d.]+) x ([\d.]+) pts \(A4\)", info)
    check(m and float(m.group(1)) > float(m.group(2)), "PDF：A4 横（%s）" % (m.group(0) if m else info))
    txt = subprocess.run(["pdftotext", "-layout", VP, "-"], stdout=subprocess.PIPE, universal_newlines=True).stdout
    hdr = next((l for l in txt.splitlines() if "項目名" in l), "")
    pos = [hdr.find(c) for c in ("項目名", "最終値", "要求値", "判定", "入力元（手順/ページ）", "入力日時")]
    check(all(p >= 0 for p in pos) and pos == sorted(pos), "PDF：列 項目名 / 最終値 / 要求値 / 判定 / 入力元（手順/ページ）/ 入力日時 の順")
    check("作業入力値一覧" in txt and "web01" in txt and "サーバ名" in txt and "作業日" in txt and "作業者" in txt and "山田 太朗" in txt and "2026-10-03" in txt, "PDF：見出し（サーバ名・作業日・作業者）")
    check("t3.medium" in txt and "不一致" in txt and "食い違い" in txt and "山田 太朗" in txt and "手順 1" in txt, "PDF：不一致・食い違い・入力元（手順）")
    fonts = subprocess.run(["pdffonts", VP], stdout=subprocess.PIPE, universal_newlines=True).stdout if shutil.which("pdffonts") else ""
    check(not fonts or re.search(r"(?i)(cjk|gothic|meiryo|jp)", fonts), "PDF：日本語フォントが埋め込まれている")
    if shutil.which("pdftoppm"):
        SF = os.path.join(ROOT, "screenshots_feature")
        os.makedirs(SF, exist_ok=True)
        for f in os.listdir(SF):
            if f.startswith("04-values-pdf"):
                os.remove(os.path.join(SF, f))
        subprocess.run(["pdftoppm", "-png", "-r", "110", VP, os.path.join(SF, "04-values-pdf")], check=True)
        print("    render: values.pdf → screenshots_feature/04-values-pdf-*.png")

print("\nRESULT: %d passed, %d failed" % (ok, bad))
sys.exit(1 if bad else 0)
