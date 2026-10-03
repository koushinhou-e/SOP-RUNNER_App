"""参数表 (.xlsx) 解析：行 = 参数，列 = 服务器（或单列“値”）。

支持的表头关键字（同一行内）:
  项目列: 項目 / 項目名 / パラメータ / 設定項目 / item / name
  键列(可选): キー / key / 変数名 / パラメータID / id
  分类列(可选): 区分 / 分類 / カテゴリ / category
  值列: 键列/项目列右侧、表头非空且不是 備考/説明 的列，每列一个服务器（表头 = 服务器名）
"""
import datetime
import re

from openpyxl import load_workbook

from .xlsx_detect import slug

ITEM_HDR = re.compile(r"^(項目|項目名|パラメータ|パラメータ名|設定項目|item|name|名称)$", re.I)
KEY_HDR = re.compile(r"^(キー|key|変数名|変数|パラメータid|id|プレースホルダ)$", re.I)
CAT_HDR = re.compile(r"^(区分|分類|カテゴリ|カテゴリー|category|大項目)$", re.I)
SKIP_HDR = re.compile(r"(備考|説明|メモ|note|remark|comment|description)", re.I)


def _s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(sep=" ")
    if isinstance(v, datetime.date):
        return v.isoformat()
    return str(v).strip()


def parse(path):
    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        res = _parse_sheet(ws)
        if res and res["params"]:
            return res
    raise ValueError("パラメータシートの見出し行が見つかりません（「項目」や「キー」などの列名を含む行が必要です）")


def _parse_sheet(ws):
    maxc = min(ws.max_column, 40)
    merged_val = {}
    for mr in ws.merged_cells.ranges:
        v = ws.cell(row=mr.min_row, column=mr.min_col).value
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                merged_val[(r, c)] = v

    def val(r, c):
        if (r, c) in merged_val:
            return merged_val[(r, c)]
        return ws.cell(row=r, column=c).value

    for hr in range(1, min(ws.max_row, 30) + 1):
        hdr = [_s(val(hr, c)) for c in range(1, maxc + 1)]
        item_c = next((i + 1 for i, h in enumerate(hdr) if ITEM_HDR.match(h)), None)
        key_c = next((i + 1 for i, h in enumerate(hdr) if KEY_HDR.match(h)), None)
        if not item_c and not key_c:
            continue
        cat_c = next((i + 1 for i, h in enumerate(hdr) if CAT_HDR.match(h)), None)
        start = max(x for x in (item_c, key_c, cat_c) if x) + 1
        servers = []
        for c in range(start, maxc + 1):
            h = hdr[c - 1]
            if h and not SKIP_HDR.search(h):
                servers.append((c, h))
        if not servers:
            continue
        if len(servers) == 1 and re.match(r"^(値|設定値|value|期待値|パラメータ値)$", servers[0][1], re.I):
            servers = [(servers[0][0], "default")]
        params, blank_run, cat = [], 0, ""
        for r in range(hr + 1, ws.max_row + 1):
            label = _s(val(r, item_c)) if item_c else ""
            key = _s(val(r, key_c)) if key_c else ""
            if cat_c and _s(val(r, cat_c)):
                cat = _s(val(r, cat_c))
            if not label and not key:
                blank_run += 1
                if blank_run >= 5:
                    break
                continue
            blank_run = 0
            params.append({
                "key": key or slug(label),
                "label": label or key,
                "category": cat,
                "row": r,
                "values": {name: _s(val(r, c)) for c, name in servers},
            })
        keys = set()
        for p in params:  # 去重键名
            k, i = p["key"], 2
            while p["key"] in keys:
                p["key"] = "%s_%d" % (k, i)
                i += 1
            keys.add(p["key"])
        return {"sheet": ws.title, "header_row": hr, "servers": [n for _, n in servers], "params": params}
    return None


def server_context(parsed, server):
    return {p["key"]: p["values"].get(server, "") for p in parsed.get("params", [])}
