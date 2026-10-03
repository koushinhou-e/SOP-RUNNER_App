"""参数比对：期待值（参数表） vs 实测值（作业中填写）。"""
import datetime
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_NUM_UNIT = re.compile(r"^([+-]?\d+(?:\.\d+)?)\s*([a-zA-Z%]*)$")
_UNIT_ALIAS = {"gib": "gib", "gb": "gib", "g": "gib", "mib": "mib", "mb": "mib", "m": "mib", "tib": "tib", "tb": "tib", "t": "tib",
               "": "", "core": "", "cores": "", "vcpu": "", "個": ""}


def normalize(s):
    s = unicodedata.normalize("NFKC", "" if s is None else str(s)).strip()
    return re.sub(r"\s+", " ", s).lower()


def judge(expected, actual):
    """→ 'match' | 'mismatch' | 'missing' | 'no_expected'"""
    e, a = normalize(expected), normalize(actual)
    if a == "":
        return "missing"
    if e == "":
        return "no_expected"
    if e == a:
        return "match"
    me, ma = _NUM_UNIT.match(e.replace(",", "")), _NUM_UNIT.match(a.replace(",", ""))
    if me and ma:
        ue, ua = _UNIT_ALIAS.get(me.group(2), me.group(2)), _UNIT_ALIAS.get(ma.group(2), ma.group(2))
        if (ue == ua or ue == "" or ua == "") and float(me.group(1)) == float(ma.group(1)):
            return "match"
        return "mismatch"
    te, ta = [t for t in re.split(r"[,\s]+", e) if t], [t for t in re.split(r"[,\s]+", a) if t]
    if len(te) > 1 or len(ta) > 1:
        return "match" if sorted(te) == sorted(ta) else "mismatch"
    return "mismatch"


STATUS_LABEL = {"match": "一致", "mismatch": "不一致", "missing": "未填写", "no_expected": "无期待值"}


def build_rows(params, server, state, hints=None):
    state = state or {}
    hints = hints or {}
    rows = []
    for p in params:
        st = state.get(p["key"]) or {}
        exp = p["values"].get(server, "")
        actual = st.get("actual", "")
        auto = judge(exp, actual)
        rows.append({
            "key": p["key"], "label": p["label"], "category": p.get("category", ""), "expected": exp, "actual": actual,
            "auto": auto, "auto_label": STATUS_LABEL[auto], "judgement": st.get("judgement") or "", "judged_at": st.get("judged_at") or "",
            "note": st.get("note", ""), "commands": hints.get(p["key"], []),
        })
    return rows


def summary(rows):
    s = {"total": len(rows), "match": 0, "mismatch": 0, "missing": 0, "no_expected": 0, "ok": 0, "ng": 0, "unjudged": 0}
    for r in rows:
        s[r["auto"]] += 1
        if r["judgement"] == "OK":
            s["ok"] += 1
        elif r["judgement"] == "NG":
            s["ng"] += 1
        else:
            s["unjudged"] += 1
    return s


def export_xlsx(rows, meta, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "パラメータ比較"
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "パラメータ比較結果"
    ws["A1"].font = Font(size=14, bold=True)
    info = [("作業名", meta.get("job_name", "")), ("対象サーバ", meta.get("server", "")), ("パラメータシート", meta.get("param_sheet", "")),
            ("作業者", meta.get("worker", "")), ("出力日時", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))]
    s = summary(rows)
    info.append(("集計", "全%d項目 / 一致%d / 不一致%d / 未入力%d / OK%d / NG%d / 未判定%d" % (
        s["total"], s["match"], s["mismatch"], s["missing"], s["ok"], s["ng"], s["unjudged"])))
    for i, (k, v) in enumerate(info):
        ws.cell(row=3 + i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=3 + i, column=2, value=v)
    hr = 4 + len(info)
    headers = ["No", "区分", "項目", "キー", "期待値", "実測値", "自動判定", "判定(OK/NG)", "判定日時", "備考"]
    widths = [5, 12, 22, 18, 28, 28, 10, 11, 19, 30]
    hdr_fill = PatternFill("solid", fgColor="D9E2F3")
    for i, h in enumerate(headers):
        c = ws.cell(row=hr, column=i + 1, value=h)
        c.font, c.fill, c.border = Font(bold=True), hdr_fill, border
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[c.column_letter].width = widths[i]
    red, yellow = PatternFill("solid", fgColor="FDE2E2"), PatternFill("solid", fgColor="FFF4CC")
    for n, r in enumerate(rows, 1):
        vals = [n, r["category"], r["label"], r["key"], r["expected"], r["actual"], r["auto_label"], r["judgement"], r["judged_at"], r["note"]]
        for i, v in enumerate(vals):
            c = ws.cell(row=hr + n, column=i + 1, value=v)
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if r["auto"] == "mismatch" or r["judgement"] == "NG":
                c.fill = red
            elif r["auto"] == "missing":
                c.fill = yellow
        if r["judgement"] == "NG":
            ws.cell(row=hr + n, column=8).font = Font(bold=True, color="C00000")
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    wb.save(out_path)
