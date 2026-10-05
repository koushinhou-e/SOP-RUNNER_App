"""把已知值写入客户交付物模板并另存（保留模板原有格式）。

只修改值：openpyxl 写入 cell.value 时保留字体/填充/边框/数字格式/合并单元格/列宽/数据验证/条件格式。
注意：openpyxl 不保留图片、图表、形状、宏（见 README 的已知限制）。
"""
import datetime
import re

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.styles.numbers import is_date_format

from .images import fit, raw_image
from .rules import to_number

_DATE = re.compile(r"^\s*(\d{4})[-/](\d{1,2})[-/](\d{1,2})\s*$")
_PLAIN_NUM = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")     # 先頭ゼロ付き（ID など）は文字列のまま


def _typed(item, value, cell):
    s = str(value).strip()
    if item.get("type") == "number" or ((item.get("source") or {}).get("kind") == "result" and _PLAIN_NUM.match(s)):
        n = to_number(s)
        if n is not None:
            return int(n) if n.is_integer() else n
    if item.get("type") == "date" or is_date_format(cell.number_format or "General"):
        m = _DATE.match(s)
        if m:
            try:
                d = datetime.datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                if is_date_format(cell.number_format or "General"):
                    return d
            except ValueError:
                pass
    return s


def fill_template(template_path, items, values, out_path, evidence=None):
    """values: {item_id: value}。evidence: 証跡画像のリスト（Service.evidence_images）。書き込んだセルの ID リストを返す。"""
    wb = load_workbook(template_path)
    written = []
    for it in items:
        if it["id"] not in values:
            continue
        v = values[it["id"]]
        if v is None or str(v).strip() == "":
            continue
        if it["sheet"] not in wb.sheetnames:
            continue
        ws = wb[it["sheet"]]
        cell = ws[it["cell"]]
        if type(cell).__name__ == "MergedCell":
            continue
        ph = it.get("placeholder")
        cur = cell.value
        if ph and isinstance(cur, str) and ph in cur and cur.strip() != ph.strip():
            cell.value = cur.replace(ph, str(v), 1)       # 保留占位符前后的文字
        else:
            cell.value = _typed(it, v, cell)
        written.append(it["id"])
    if evidence:
        written += add_evidence(wb, items, evidence)
    wb.save(out_path)
    return written


def _read(path):
    with open(path, "rb") as f:
        return f.read()


def _unique_sheet_name(wb, base):
    name, n = base, 2
    while name in wb.sheetnames:
        name = "%s%d" % (base, n)
        n += 1
    return name


def _col_px(ws, letter):
    w = ws.column_dimensions[letter].width if letter in ws.column_dimensions else None
    return int((w or 8.43) * 7 + 5)


def _row_px(ws, r):
    h = ws.row_dimensions[r].height if r in ws.row_dimensions else None
    return int((h or 15) * 4 / 3)


def _cell_box(ws, coord):
    """セル（結合セルなら結合範囲全体）の大きさ（px の概算）。"""
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import coordinate_to_tuple
    r, c = coordinate_to_tuple(coord)
    r1, c1, r2, c2 = r, c, r, c
    for rg in ws.merged_cells.ranges:
        if rg.min_row <= r <= rg.max_row and rg.min_col <= c <= rg.max_col:
            r1, c1, r2, c2 = rg.min_row, rg.min_col, rg.max_row, rg.max_col
            break
    w = sum(_col_px(ws, get_column_letter(x)) for x in range(c1, c2 + 1))
    h = sum(_row_px(ws, y) for y in range(r1, r2 + 1))
    return "%s%d" % (get_column_letter(c1), r1), w, h


def add_evidence(wb, items, evidence, sheet_title="証跡", max_w=620, max_h=820, cell_max_w=360, cell_max_h=270):
    """証跡画像を挿入する（Pillow 不要）。

    * マッピングで「証跡画像」（source.kind = evidence, key = 証跡キー）が指定されたセル → そのキーの最初の画像をセル位置に配置
      （結合セルなど十分な大きさがあればその範囲に収まるよう縮小、小さいセルなら最大 cell_max_w×cell_max_h で左上に配置）
    * すべての画像を末尾の「証跡」シートに、手順番号・手順名・説明つきで縦に並べて配置
    """
    placed = []
    by_key = {}
    for e in evidence:
        if e[3] and e[3] not in by_key:
            by_key[e[3]] = e
    for it in items:
        src = it.get("source") or {}
        if src.get("kind") != "evidence" or it.get("sheet") not in wb.sheetnames:
            continue
        e = by_key.get(src.get("key"))
        if not e:
            continue
        meta = e[5]
        ws_t = wb[it["sheet"]]
        anchor, bw, bh = _cell_box(ws_t, it["cell"])
        mw, mh = (bw - 4, bh - 4) if bw >= 120 and bh >= 80 else (cell_max_w, cell_max_h)
        w, h = fit(meta.get("w") or 1, meta.get("h") or 1, mw, mh)
        img = raw_image(_read(e[4]), w, h)
        img.anchor = anchor
        ws_t.add_image(img)
        placed.append(it["id"])
    ws = wb.create_sheet(_unique_sheet_name(wb, sheet_title))
    ws.column_dimensions["A"].width = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True      # 印刷時は幅 1 ページに収める
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws["A1"] = "証跡画像"
    ws["A1"].font = Font(bold=True, size=14)
    row = 3
    for no, title, desc, key, path, meta in evidence:
        ws.cell(row=row, column=1, value="手順 %d：%s" % (no, title)).font = Font(bold=True)
        ws.cell(row=row + 1, column=1, value=(desc or "（説明なし）") + ("　[%s]" % key if key else ""))
        w, h = fit(meta.get("w") or 1, meta.get("h") or 1, max_w, max_h)
        img = raw_image(_read(path), w, h)
        ws.row_dimensions[row].height = ws.row_dimensions[row + 1].height = 18
        img.anchor = "A%d" % (row + 3)
        ws.add_image(img)
        row += 3 + int(h / 20) + 3          # 既定の行の高さ 15pt ≒ 20px
    return placed
