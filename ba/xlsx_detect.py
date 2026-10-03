"""Excel 交付物模板：检测需要填写的单元格。

检测来源 (item["reason"]):
  placeholder  单元格文字含 {{name}} / ＿＿＿ / ___ / （　） / 【　】
  validation   单元格带“数据验证”(列表 / 整数 / 小数)
  highlight    单元格为空且有底色（实心填充，非白色）
  label        单元格为空、带边框（或左侧标签以：结尾），且左侧有文字标签
标签 = 左侧最近的文字标签（行标签）＋ 上方最近的表头（列标题，若不是“設定値”之类的通用词）。
"""
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z_][\w.-]*)\s*\}\}|[＿_]{3,}|＿{2,}|[（(][ 　]*[）)]|【[ 　]*】")
GENERIC_HEADER = re.compile(r"^(設定値|値|内容|入力|入力欄|記入欄|記入|value|values|パラメータ値|実績値|設定内容|値/内容)$", re.I)
OPTIONAL_HEADER = re.compile(r"(備考|メモ(?!リ)|コメント|特記|note|remark|comment)", re.I)
NUM_LABEL = re.compile(r"(数|サイズ|容量|GiB|GB|MiB|MB|TB|vCPU|CPU|台数|件数|ポート|port|size|count|メモリ)", re.I)
DATE_LABEL = re.compile(r"(日付|作業日|実施日|確認日|date)", re.I)
PRESET_BY_LABEL = [
    (re.compile(r"(IP|ＩＰ)", re.I), "ipv4"),
    (re.compile(r"(インスタンスID|instance.?id)", re.I), "instance_id"),
    (re.compile(r"AMI", re.I), "ami_id"),
    (re.compile(r"(サブネット|subnet)", re.I), "subnet_id"),
    (re.compile(r"(セキュリティグループ|security.?group|^SG$)", re.I), "sg_id"),
    (re.compile(r"VPC", re.I), "vpc_id"),
    (re.compile(r"(インスタンスタイプ|instance.?type)", re.I), "instance_type"),
    (re.compile(r"(ホスト名|サーバ名|サーバー名|hostname|FQDN)", re.I), "hostname"),
]
MAX_ROWS, MAX_COLS, MAX_DV_CELLS = 400, 60, 5000


def is_empty(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def norm_label(s):
    s = unicodedata.normalize("NFKC", str(s or "")).strip().lower()
    return re.sub(r"[\s:：]+", "", s)


def base_label(s):
    return re.sub(r"[\(（\[［【].*?[\)）\]］】]", "", norm_label(s))


def _highlighted(cell):
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return False
    c = f.fgColor
    if c is None:
        return False
    if c.type == "rgb":
        rgb = (c.rgb or "").upper() if isinstance(c.rgb, str) else ""
        return rgb not in ("", "FFFFFFFF", "00FFFFFF", "00000000")
    if c.type == "theme":
        return not (c.theme in (0,) and not c.tint)
    if c.type == "indexed":
        return c.indexed not in (9, 64, 65)
    return False


def _has_border(cell):
    b = cell.border
    return any(getattr(getattr(b, side), "style", None) for side in ("left", "right", "top", "bottom"))


def _sheet_values(wb, ref, ws):
    """解析数据验证列表引用（A1:A5 / 'シート'!$A$1:$A$5 / 名前定义）→ 值列表。"""
    ref = ref.strip().lstrip("=")
    targets = []
    if "!" in ref:
        sh, rng = ref.rsplit("!", 1)
        targets.append((sh.strip("'"), rng))
    elif re.match(r"^\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?$", ref):
        targets.append((ws.title, ref))
    else:  # defined name
        dn = None
        try:
            dn = wb.defined_names.get(ref)
        except Exception:
            dn = None
        if dn is None:
            dn = ws.defined_names.get(ref) if hasattr(ws, "defined_names") else None
        if dn is not None:
            for sh, rng in dn.destinations:
                targets.append((sh, rng))
    out = []
    for sh, rng in targets:
        if sh not in wb.sheetnames:
            continue
        s = wb[sh]
        mc, mr, xc, xr = range_boundaries(rng.replace("$", ""))
        for row in s.iter_rows(min_row=mr, max_row=min(xr, mr + 500), min_col=mc, max_col=xc):
            for c in row:
                if not is_empty(c.value):
                    out.append(str(c.value).strip())
    return out


def _dv_info(dv, wb, ws):
    t = dv.type
    f1 = (dv.formula1 or "").strip()
    if t == "list":
        if f1.startswith('"') and f1.endswith('"'):
            opts = [o.strip() for o in f1[1:-1].split(",") if o.strip()]
        else:
            opts = _sheet_values(wb, f1, ws)
        return {"type": "dropdown", "options": opts}
    if t in ("whole", "decimal"):
        info = {"type": "number", "options": []}
        op = dv.operator or "between"

        def num(x):
            try:
                return float(x) if t == "decimal" else int(float(x))
            except (TypeError, ValueError):
                return None
        a, b = num(dv.formula1), num(dv.formula2)
        if op == "between":
            info["min"], info["max"] = a, b
        elif op in ("greaterThan", "greaterThanOrEqual"):
            info["min"] = a
        elif op in ("lessThan", "lessThanOrEqual"):
            info["max"] = a
        return info
    if t == "date":
        return {"type": "date", "options": []}
    return None


WELL_KNOWN_KEYS = {"作業日": "work_date", "実施日": "work_date", "作業者": "worker", "実施者": "worker", "確認者": "reviewer",
                   "承認者": "approver", "案件名": "project_name", "プロジェクト名": "project_name"}


def slug(s):
    s = unicodedata.normalize("NFKC", str(s or "")).strip()
    if s in WELL_KNOWN_KEYS:
        return WELL_KNOWN_KEYS[s]
    if re.search(r"[^\x00-\x7f]", s):          # 含日文/中文：保留原文作为键
        return re.sub(r"\s+", "_", s)
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_").lower() or s


def detect_workbook(path):
    wb = load_workbook(path)
    items = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        items.extend(_detect_sheet(wb, ws))
    seen = set()
    out = []
    for it in items:
        if it["id"] in seen:
            continue
        seen.add(it["id"])
        out.append(it)
    return out


def _detect_sheet(wb, ws):
    covered = {}   # (r,c) -> top-left (r,c)
    tl_span = {}   # top-left -> (max_r, max_c)
    for mr in ws.merged_cells.ranges:
        tl_span[(mr.min_row, mr.min_col)] = (mr.max_row, mr.max_col)
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                if (r, c) != (mr.min_row, mr.min_col):
                    covered[(r, c)] = (mr.min_row, mr.min_col)
    dv_map = {}
    for dv in ws.data_validations.dataValidation:
        info = _dv_info(dv, wb, ws)
        if not info:
            continue
        n = 0
        for rng in dv.sqref.ranges:
            for r in range(rng.min_row, min(rng.max_row, MAX_ROWS) + 1):
                for c in range(rng.min_col, min(rng.max_col, MAX_COLS) + 1):
                    if n < MAX_DV_CELLS:
                        dv_map[(r, c)] = info
                        n += 1
    maxr, maxc = min(ws.max_row, MAX_ROWS), min(ws.max_column, MAX_COLS)

    def val(r, c):
        r, c = covered.get((r, c), (r, c))
        return ws.cell(row=r, column=c).value

    def is_label_text(v):
        return isinstance(v, str) and v.strip() and not PLACEHOLDER.search(v) and len(v.strip()) <= 40

    def left_label(r, c):
        cc = c - 1
        steps = 0
        while cc >= 1 and steps < 6:
            if (r, cc) in covered and covered[(r, cc)][0] != r:
                break
            v = val(r, cc)
            if not is_empty(v):
                return (v.strip(), c - cc) if is_label_text(v) else (None, 0)
            cc -= 1
            steps += 1
        return (None, 0)

    def header_above(r, c):
        rr = r - 1
        steps = 0
        while rr >= 1 and steps < 40:
            v = val(rr, c)
            if not is_empty(v):
                if is_label_text(v) and (rr, c) not in dv_map:
                    # 是否是“表头”：同一行至少 2 个文字单元格
                    tops = set()
                    for cc in range(1, maxc + 1):
                        if is_label_text(val(rr, cc)):
                            tops.add(covered.get((rr, cc), (rr, cc)))
                    if len(tops) >= 2:
                        return v.strip()
                    return None
                if (rr, c) in dv_map or ws.cell(row=rr, column=c).value is None:
                    pass
                else:
                    return None
            rr -= 1
            steps += 1
        return None

    items = []
    for r in range(1, maxr + 1):
        for c in range(1, maxc + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            v = cell.value
            reason, placeholder, ph_name = None, None, None
            dv = dv_map.get((r, c))
            if isinstance(v, str) and PLACEHOLDER.search(v):
                m = PLACEHOLDER.search(v)
                reason, placeholder, ph_name = "placeholder", m.group(0), m.group(1)
            elif dv is not None:
                reason = "validation"
            elif is_empty(v) and _highlighted(cell):
                reason = "highlight"
            elif is_empty(v):
                lab, dist = left_label(r, c)
                if lab and (_has_border(cell) or (dist == 1 and re.search(r"[:：]$", lab))):
                    reason = "label"
            if not reason:
                continue
            row_label, _ = left_label(r, c)
            if reason == "placeholder":
                before = v[: v.find(placeholder)].strip()
                before = re.sub(r"[:：\s]+$", "", before)
                if before:
                    row_label = before
            col_header = header_above(r, c)
            if col_header and row_label and norm_label(col_header) == norm_label(row_label):
                col_header = None
            generic = bool(col_header and GENERIC_HEADER.match(norm_label(col_header)))
            if generic:
                col_header = None
            parts = [p for p in [row_label and re.sub(r"[:：\s]+$", "", row_label), col_header] if p]
            label = " / ".join(parts) if parts else (ph_name or "%s%d" % (get_column_letter(c), r))
            itype = "text"
            options, rules = [], {"required": True}
            if dv:
                itype = dv["type"]
                options = dv.get("options", [])
                for k in ("min", "max"):
                    if dv.get(k) is not None:
                        rules[k] = dv[k]
            else:
                nf = cell.number_format or "General"
                if is_date_format(nf) or DATE_LABEL.search(label):
                    itype = "date"
                elif nf not in ("General", "@") and re.search(r"[0#]", nf):
                    itype = "number"
                elif NUM_LABEL.search(row_label or label) and not col_header:
                    itype = "number"
            if itype == "date":
                rules["preset"] = "date"
            elif itype == "text" and not col_header:
                for rx, preset in PRESET_BY_LABEL:
                    if rx.search(row_label or label or ""):
                        rules["preset"] = preset
                        break
            col_optional = bool(OPTIONAL_HEADER.search(col_header or ""))
            optional = col_optional or bool(OPTIONAL_HEADER.search(label))
            if optional:
                rules["required"] = False
            items.append({
                "id": "%s!%s" % (ws.title, cell.coordinate),
                "sheet": ws.title,
                "cell": cell.coordinate,
                "label": label,
                "row_label": row_label,
                "col_header": col_header,
                "key": ph_name or (slug(row_label) if row_label and not col_header else None),
                "type": itype,
                "options": options,
                "default": None if is_empty(v) or reason == "placeholder" else (v if isinstance(v, (int, float)) else str(v)),
                "placeholder": placeholder,
                "reason": reason,
                "rules": rules,
                "source": {"kind": "none" if col_optional else "input"},
            })
    return items


def auto_map(items, params):
    """根据参数表自动把模板单元格映射到参数键（只映射纯“行标签”项，不映射 xx / 確認結果 这类列）。"""
    by_key = {norm_label(p["key"]): p["key"] for p in params}
    by_label, by_base = {}, {}
    for p in params:
        by_label.setdefault(norm_label(p["label"]), p["key"])
        by_base.setdefault(base_label(p["label"]), p["key"])
    n = 0
    for it in items:
        if it.get("col_header"):
            continue
        if (it.get("source") or {}).get("kind") not in (None, "input", "none") or it.get("_manual"):
            continue
        cands = [it.get("key"), it.get("row_label"), it.get("label")]
        hit = None
        for cnd in cands:
            if not cnd:
                continue
            k = norm_label(cnd)
            hit = by_key.get(k) or by_label.get(k) or by_base.get(base_label(cnd))
            if hit:
                break
        if hit:
            it["source"] = {"kind": "param", "key": hit}
            n += 1
    return n


def preview(path, max_rows=60, max_cols=16):
    """给映射界面用的网格预览。"""
    wb = load_workbook(path)
    sheets = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        maxr, maxc = min(ws.max_row, max_rows), min(ws.max_column, max_cols)
        merges = []
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= maxr and mr.min_col <= maxc:
                merges.append([mr.min_row, mr.min_col, min(mr.max_row, maxr), min(mr.max_col, maxc)])
        rows = []
        for r in range(1, maxr + 1):
            row = []
            for c in range(1, maxc + 1):
                v = ws.cell(row=r, column=c).value
                row.append("" if v is None else str(v)[:60])
            rows.append(row)
        sheets.append({"name": ws.title, "rows": rows, "merges": merges, "max_row": ws.max_row, "max_col": ws.max_column,
                       "cols": [get_column_letter(c) for c in range(1, maxc + 1)]})
    return sheets
