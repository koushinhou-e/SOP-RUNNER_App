# -*- coding: utf-8 -*-
"""生成虚构的示例数据（全部为假数据：example.local / 192.0.2.x / 198.51.100.x / 假 ID）。

    python samples/make_samples.py
"""
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "vendor"))
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

thin = Side(style="thin", color="000000")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill("solid", fgColor="DDEBF7")
YELLOW = PatternFill("solid", fgColor="FFFF99")

SERVERS = ["web01", "web02", "db01"]
PARAMS = [
    ("基本", "サーバ名", "hostname", ["web01.example.local", "web02.example.local", "db01.example.local"]),
    ("基本", "インスタンスID", "instance_id", ["i-0123456789abcdef0", "i-0123456789abcdef1", "i-0fedcba9876543210"]),
    ("基本", "インスタンスタイプ", "instance_type", ["t3.large", "t3.large", "r5.xlarge"]),
    ("基本", "vCPU", "vcpu", [2, 2, 4]),
    ("基本", "メモリ(GiB)", "memory_gib", [8, 8, 32]),
    ("OS", "OS", "os", ["Amazon Linux 2023", "Amazon Linux 2023", "Red Hat Enterprise Linux 9"]),
    ("OS", "AMI ID", "ami_id", ["ami-0abcdef1234567890", "ami-0abcdef1234567890", "ami-0fedcba0987654321"]),
    ("ネットワーク", "サブネットID", "subnet_id", ["subnet-0123456789abcdef0", "subnet-0123456789abcdef1", "subnet-0aaaabbbbccccdddd"]),
    ("ネットワーク", "セキュリティグループ", "security_group_id", ["sg-0123456789abcdef0", "sg-0123456789abcdef0", "sg-0fedcba9876543210"]),
    ("ネットワーク", "プライベートIP", "private_ip", ["192.0.2.11", "192.0.2.12", "198.51.100.21"]),
    ("ストレージ", "EBSサイズ(GiB)", "ebs_size_gib", [30, 30, 100]),
    ("共通", "リージョン", "region", ["ap-northeast-1"] * 3),
    ("共通", "SSHユーザ", "ssh_user", ["ec2-user", "ec2-user", "ec2-user"]),
]


def param_sheet(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "パラメータ"
    ws["A1"] = "EC2 パラメータシート（サンプル・架空データ）"
    ws["A1"].font = Font(size=14, bold=True)
    hdr = ["区分", "項目", "キー"] + SERVERS + ["備考"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=3, column=i, value=h)
        c.font, c.fill, c.border = Font(bold=True), HDR, BORDER
    for r, (cat, label, key, vals) in enumerate(PARAMS, 4):
        row = [cat, label, key] + vals + [""]
        for i, v in enumerate(row, 1):
            ws.cell(row=r, column=i, value=v).border = BORDER
    for col, w in zip("ABCDEFG", [12, 20, 18, 26, 26, 28, 16]):
        ws.column_dimensions[col].width = w
    wb.save(path)


def template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "構築結果"
    lists = wb.create_sheet("リスト")
    lists["A1"], lists["A2"] = "合格", "不合格"
    lists.sheet_state = "hidden"
    ws.merge_cells("B1:E1")
    ws["B1"] = "EC2サーバ構築 確認結果報告書"
    ws["B1"].font = Font(size=16, bold=True)
    ws["B1"].alignment = Alignment(horizontal="center")
    for col, w in zip("ABCDE", [3, 22, 30, 12, 28]):
        ws.column_dimensions[col].width = w
    meta = [(3, "案件名", "{{project_name}}"), (4, "作業日", None), (5, "作業者", None), (6, "確認者", "＿＿＿＿＿＿")]
    for r, lab, v in meta:
        ws.cell(row=r, column=2, value=lab).font = Font(bold=True)
        c = ws.cell(row=r, column=3, value=v)
        c.border = BORDER
        if v is None:
            c.fill = YELLOW
    ws["C4"].number_format = "yyyy/mm/dd"
    for i, h in enumerate(["項目", "設定値", "確認結果", "備考"], 2):
        c = ws.cell(row=8, column=i, value=h)
        c.font, c.fill, c.border = Font(bold=True), HDR, BORDER
        c.alignment = Alignment(horizontal="center")
    rows = ["サーバ名", "インスタンスID", "インスタンスタイプ", "vCPU", "メモリ(GiB)", "OS", "AMI ID", "サブネットID",
            "セキュリティグループ", "プライベートIP", "EBSサイズ(GiB)"]
    dv = DataValidation(type="list", formula1='"OK,NG,対象外"', allow_blank=True)
    ws.add_data_validation(dv)
    for r, lab in enumerate(rows, 9):
        ws.cell(row=r, column=2, value=lab).border = BORDER
        for c in (3, 4, 5):
            ws.cell(row=r, column=c).border = BORDER
        if lab in ("vCPU", "メモリ(GiB)", "EBSサイズ(GiB)"):
            ws.cell(row=r, column=3).number_format = "0"
        dv.add("D%d" % r)
    vcpu = DataValidation(type="whole", operator="between", formula1="1", formula2="128")
    ws.add_data_validation(vcpu)
    vcpu.add("C12")
    r = 9 + len(rows) + 1
    ws.cell(row=r, column=2, value="総合判定").font = Font(bold=True)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    ws.cell(row=r, column=3).border = BORDER
    dv2 = DataValidation(type="list", formula1="=リスト!$A$1:$A$2", allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("C%d" % r)
    ws.cell(row=r + 1, column=2, value="特記事項").font = Font(bold=True)
    ws.cell(row=r + 1, column=3, value="（　　　）")
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A9"
    wb.save(path)


COMMANDS = {
    "name": "EC2 構築確認（読み取り専用）サンプル",
    "templates": [
        {"title": "インスタンス基本情報", "kind": "aws", "checks": "instance_type,private_ip,ami_id,subnet_id",
         "template": "aws ec2 describe-instances --region {{region}} --instance-ids {{instance_id}} --query 'Reservations[].Instances[].[InstanceId,InstanceType,State.Name,PrivateIpAddress,ImageId,SubnetId]' --output text --no-cli-pager"},
        {"title": "セキュリティグループ", "kind": "aws", "checks": "security_group_id",
         "template": "aws ec2 describe-instances --region {{region}} --instance-ids {{instance_id}} --query 'Reservations[].Instances[].SecurityGroups[].GroupId' --output text --no-cli-pager"},
        {"title": "EBS ボリューム", "kind": "aws", "checks": "ebs_size_gib",
         "template": "aws ec2 describe-volumes --region {{region}} --filters Name=attachment.instance-id,Values={{instance_id}} --query 'Volumes[].[VolumeId,Size,VolumeType,State]' --output text --no-cli-pager"},
        {"title": "AMI 名", "kind": "aws", "checks": "os",
         "template": "aws ec2 describe-images --region {{region}} --image-ids {{ami_id}} --query 'Images[].[ImageId,Name]' --output text --no-cli-pager"},
        {"title": "ホスト名", "kind": "linux", "checks": "hostname",
         "template": "ssh -n -o BatchMode=yes -o ConnectTimeout=10 {{ssh_user}}@{{private_ip}} 'hostname -f'"},
        {"title": "メモリ (GiB)", "kind": "linux", "checks": "memory_gib",
         "template": "ssh -n -o BatchMode=yes -o ConnectTimeout=10 {{ssh_user}}@{{private_ip}} 'free -g'"},
        {"title": "vCPU 数", "kind": "linux", "checks": "vcpu",
         "template": "ssh -n -o BatchMode=yes -o ConnectTimeout=10 {{ssh_user}}@{{private_ip}} 'nproc'"},
        {"title": "ディスク使用量", "kind": "linux", "checks": "ebs_size_gib",
         "template": "ssh -n -o BatchMode=yes -o ConnectTimeout=10 {{ssh_user}}@{{private_ip}} 'df -h /'"},
        {"title": "OS バージョン", "kind": "linux", "checks": "os",
         "template": "ssh -n -o BatchMode=yes -o ConnectTimeout=10 {{ssh_user}}@{{private_ip}} 'grep PRETTY_NAME /etc/os-release'"},
    ],
}

# ---------------- EC2 構築確認手順書（.docx）----------------
# パラメータシート（web01 列）・成果物テンプレート・コマンド集と同じ項目を確認する手順書。
# 入力欄のラベルは PROCEDURE_KEYS でパラメータキーに対応付け、サンプル読み込み時にライブラリへ設定として保存する
# （手順実行で開くとキーが自動で付き、④ 最終値チェック・主要値一覧 PDF で ③ の実測値と突き合わせられる）。
PROCEDURE_NAME = "EC2構築確認手順書_sample.docx"
PROCEDURE_BASE = "Webサーバ定期パッチ適用手順書.docx"      # スタイル・番号付けを流用する
BL = "＿＿＿＿＿＿＿＿"
PROCEDURE = [
    ("title", "EC2サーバ構築確認手順書（web01）"),
    ("p", "文書番号：OPS-EC2-0001（架空）　版数：1.0　※本書はサンプル用の架空の手順書です。"),
    ("table", ["作業日", "作業者", "確認者"]),
    ("h1", "1. 事前準備"),
    ("step", "作業端末から踏み台サーバ（192.0.2.10）へSSH接続する。"),
    ("code", "$ ssh ec2-user@192.0.2.10"),
    ("p", "期待結果：踏み台サーバのプロンプトが表示されること。"),
    ("h1", "2. インスタンス情報の確認"),
    ("step", "対象インスタンスの基本情報を確認する。"),
    ("code", "$ aws ec2 describe-instances --region ap-northeast-1 --instance-ids i-0123456789abcdef0 --query 'Reservations[].Instances[].[InstanceId,InstanceType,PrivateIpAddress]' --output text --no-cli-pager"),
    ("p", "インスタンスID：" + BL),
    ("p", "インスタンスタイプ：" + BL),
    ("p", "プライベートIP：" + BL),
    ("step", "AMI とサブネットを確認する。"),
    ("code", "$ aws ec2 describe-instances --region ap-northeast-1 --instance-ids i-0123456789abcdef0 --query 'Reservations[].Instances[].[ImageId,SubnetId]' --output text --no-cli-pager"),
    ("p", "AMI ID：" + BL),
    ("p", "サブネットID：" + BL),
    ("step", "セキュリティグループを確認する。"),
    ("code", "$ aws ec2 describe-instances --region ap-northeast-1 --instance-ids i-0123456789abcdef0 --query 'Reservations[].Instances[].SecurityGroups[].GroupId' --output text --no-cli-pager"),
    ("p", "セキュリティグループ：" + BL),
    ("step", "EC2 コンソールの詳細画面で、インスタンスの状態が running であることを確認する。"),
    ("p", "確認結果：（　　）　※running なら「OK」と記入"),
    ("h1", "3. OS の確認"),
    ("step", "サーバへSSH接続し、ホスト名を確認する。"),
    ("code", "$ ssh ec2-user@192.0.2.11"),
    ("code", "$ hostname -f"),
    ("p", "ホスト名：" + BL),
    ("step", "OS のバージョンを確認する。"),
    ("code", "$ grep PRETTY_NAME /etc/os-release"),
    ("p", "OS：" + BL),
    ("step", "vCPU 数とメモリ容量を確認する。"),
    ("code", "$ nproc"),
    ("code", "$ free -g"),
    ("p", "vCPU：（　　）"),
    ("p", "メモリ：（　　）GiB"),
    ("step", "ルートボリュームのサイズを確認する。"),
    ("code", "$ lsblk -d -o NAME,SIZE"),
    ("p", "EBSサイズ：（　　）GiB"),
    ("h1", "4. 完了報告"),
    ("step", "作業結果を作業責任者へ報告する。"),
    ("p", "□ 作業責任者へ報告した"),
]
# 手順タイトル → [(入力欄のラベル, パラメータキー)]、証跡画像の要求（説明, キー）
PROCEDURE_KEYS = {
    "記録: 作業日・作業者・確認者": [("作業日", "work_date"), ("作業者", "worker")],
    "対象インスタンスの基本情報を確認する。": [("インスタンスID", "instance_id"), ("インスタンスタイプ", "instance_type"), ("プライベートIP", "private_ip")],
    "AMI とサブネットを確認する。": [("AMI ID", "ami_id"), ("サブネットID", "subnet_id")],
    "セキュリティグループを確認する。": [("セキュリティグループ", "security_group_id")],
    "サーバへSSH接続し、ホスト名を確認する。": [("ホスト名", "hostname")],
    "OS のバージョンを確認する。": [("OS", "os")],
    "vCPU 数とメモリ容量を確認する。": [("vCPU", "vcpu"), ("メモリ", "memory_gib")],
    "ルートボリュームのサイズを確認する。": [("EBSサイズ", "ebs_size_gib")],
}
PROCEDURE_EVIDENCE = {
    "EC2 コンソールの詳細画面で、インスタンスの状態が running であることを確認する。": [("EC2 詳細画面のスクリーンショット", "img_ec2")],
}


def procedure_settings():
    """ライブラリの手順書に保存する設定（web/sop/sop-evidence.js の settingsOf() と同じ形）。手順はタイトルで対応付く。"""
    titles = [t for kind, t in PROCEDURE if kind == "step"]
    order = ["記録: 作業日・作業者・確認者"] + titles
    steps = []
    for title in order:
        keys, ev = PROCEDURE_KEYS.get(title, []), PROCEDURE_EVIDENCE.get(title, [])
        if keys or ev:
            steps.append({"index": order.index(title), "title": title,
                          "items": [{"id": "ev_%s" % k, "desc": d, "key": k} for d, k in ev],
                          "inputKeys": [{"label": lb, "key": k} for lb, k in keys]})
    return {"docName": PROCEDURE_NAME, "steps": steps}


def procedure(path):
    import zipfile
    from xml.sax.saxutils import escape

    def run(t, mono=False):
        rpr = '<w:rPr><w:rFonts w:ascii="Courier New" w:hAnsi="Courier New"/><w:sz w:val="19"/></w:rPr>' if mono else ""
        return '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r>' % (rpr, escape(t))

    def para(t, style=None, indent=False, mono=False):
        ppr = ('<w:pStyle w:val="%s"/>' % style if style else "") + ('<w:ind w:left="480"/>' if indent else "")
        return "<w:p>%s%s</w:p>" % ("<w:pPr>%s</w:pPr>" % ppr if ppr else "", run(t, mono))

    def cell(t):
        return '<w:tc><w:tcPr><w:tcW w:type="dxa" w:w="4320"/></w:tcPr>%s</w:tc>' % (para(t) if t else "<w:p/>")

    body = []
    for kind, v in PROCEDURE:
        if kind == "title":
            body.append(para(v, "Title"))
        elif kind == "h1":
            body.append(para(v, "Heading1"))
        elif kind == "step":
            body.append(para(v, "ListNumber"))
        elif kind == "code":
            body.append(para(v, indent=True, mono=True))
        elif kind == "table":
            body.append('<w:tbl><w:tblPr><w:tblStyle w:val="TableGrid"/><w:tblW w:type="auto" w:w="0"/></w:tblPr>'
                        '<w:tblGrid><w:gridCol w:w="4320"/><w:gridCol w:w="4320"/></w:tblGrid>%s</w:tbl>'
                        % "".join("<w:tr>%s%s</w:tr>" % (cell(lb), cell("")) for lb in v))
        else:
            body.append(para(v, indent=True))
    with zipfile.ZipFile(os.path.join(HERE, PROCEDURE_BASE)) as src:
        doc = src.read("word/document.xml").decode("utf-8")
        head = doc[: doc.index("<w:body>") + len("<w:body>")]
        sect = doc[doc.rindex("<w:sectPr"):] if "<w:sectPr" in doc else "</w:body></w:document>"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                if item.filename == "docProps/thumbnail.jpeg":
                    continue
                data = src.read(item.filename)
                if item.filename == "word/document.xml":
                    data = (head + "".join(body) + sect).encode("utf-8")
                elif item.filename == "docProps/core.xml":
                    data = re.sub(r"<dc:title>.*?</dc:title>", "<dc:title>EC2サーバ構築確認手順書</dc:title>", data.decode("utf-8")).encode("utf-8")
                elif item.filename in ("[Content_Types].xml", "_rels/.rels"):
                    data = data.decode("utf-8")
                    data = re.sub(r'<Relationship [^>]*thumbnail[^>]*/>', "", data)
                    data = data.encode("utf-8")
                dst.writestr(item.filename, data)


if __name__ == "__main__":
    param_sheet(os.path.join(HERE, "EC2パラメータシート_sample.xlsx"))
    template(os.path.join(HERE, "構築結果報告書_template_sample.xlsx"))
    with open(os.path.join(HERE, "command_templates_sample.json"), "w", encoding="utf-8") as f:
        json.dump(COMMANDS, f, ensure_ascii=False, indent=2)
    procedure(os.path.join(HERE, PROCEDURE_NAME))
    with open(os.path.join(HERE, "EC2構築確認手順書_settings.json"), "w", encoding="utf-8") as f:
        json.dump(procedure_settings(), f, ensure_ascii=False, indent=2)
    print("samples written to", HERE)
